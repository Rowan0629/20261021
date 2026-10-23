import io
import zipfile

import pytest
from openpyxl import load_workbook
from PIL import Image

from troupe_bot.store import COLUMNS, AlreadyLoaned, AssetRow, ExcelBusy, NotLoaned, Store


def thumb_bytes(color=(10, 20, 30)) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (60, 40), color).save(buf, format="JPEG")
    return buf.getvalue()


def make_row(**over) -> AssetRow:
    base = dict(
        registered_at="2026-10-22 14:05:09",
        uploader="小明",
        name="青花瓷花瓶",
        category="小道具",
        features="青花瓷",
        scene="古典室內景",
        play="《白蛇傳》",
        quantity=1,
        condition="良好",
        location="B區03箱",
        photo_paths=["C:/data/photos/a.jpg"],
        updated_by="小明",
        updated_at="2026-10-22 14:05:09",
    )
    base.update(over)
    return AssetRow(**base)


@pytest.fixture
def store(tmp_path):
    warnings: list[str] = []
    s = Store(tmp_path / "inventory.xlsx", retries=3, delay=0.0, sleep=lambda _s: None, warn=warnings.append)
    s.warnings = warnings
    s.ensure()
    return s


# ------------------------------------------------------------ 建檔與追加


def test_new_workbook_has_all_headers_and_hides_the_message_id_column(store):
    ws = load_workbook(store.path)["資產"]
    assert [c.value for c in ws[1]] == COLUMNS
    assert len(COLUMNS) == 22
    assert COLUMNS[18:] == ["數量來源", "入庫方式", "借出人", "借出時間"]
    assert ws.column_dimensions["R"].hidden is True


def test_rows_are_appended_at_the_end_with_increasing_ids(store):
    a = store.add(make_row(name="第一件"), thumb_bytes())
    b = store.add(make_row(name="第二件"), thumb_bytes())
    assert (a.id, b.id) == (1, 2)
    ws = load_workbook(store.path)["資產"]
    assert [ws.cell(r, 4).value for r in (2, 3)] == ["第一件", "第二件"]
    assert ws.max_row == 3


def test_ids_are_never_reused_after_a_deletion(store):
    store.add(make_row(), None)
    store.mark_deleted(1, by="阿華", at="2026-10-22 15:00:00")
    assert store.add(make_row(), None).id == 2


def test_thumbnails_survive_reloading_the_file_for_the_next_row(store):
    for i in range(3):
        store.add(make_row(name=f"件{i}"), thumb_bytes((i * 50, 0, 0)))
    ws = load_workbook(store.path)["資產"]
    assert len(ws._images) == 3


def test_photo_path_cell_lists_every_photo_and_links_the_first(store):
    row = make_row(photo_paths=["C:/data/photos/a.jpg", "C:/data/photos/a-2.jpg"])
    store.add(row, None)
    cell = load_workbook(store.path)["資產"].cell(2, 14)
    assert cell.value == "C:/data/photos/a.jpg\nC:/data/photos/a-2.jpg"
    assert cell.hyperlink is not None and cell.hyperlink.target.startswith("file:///")


def test_rows_round_trip_through_the_file(store):
    store.add(make_row(condition_ai=True, remarks="借來的", quantity=3), None)
    [row] = store.all_rows()
    assert (row.quantity, row.condition, row.condition_ai, row.remarks) == (3, "良好", True, "借來的")
    assert row.photo_paths == ["C:/data/photos/a.jpg"]
    assert row.status == "正常"


# ------------------------------------------------------------ 防護 4：公式


@pytest.mark.parametrize("evil", ["=1+1", '=HYPERLINK("http://x","y")', "+1+1", "-2+3", "@SUM(A1)"])
def test_text_starting_with_formula_characters_is_stored_as_plain_text(store, evil):
    store.add(make_row(name=evil, uploader=evil, features=evil, remarks=evil), None)
    with zipfile.ZipFile(store.path) as z:
        sheet = z.read("xl/worksheets/sheet1.xml").decode("utf-8")
    assert "<f>" not in sheet and "<f " not in sheet
    ws = load_workbook(store.path)["資產"]
    for col in (3, 4, 6, 12):
        cell = ws.cell(2, col)
        assert cell.value == evil and cell.data_type == "s"
        assert cell.quotePrefix is True


def test_ordinary_text_is_not_marked_with_a_quote_prefix(store):
    store.add(make_row(), None)
    assert load_workbook(store.path)["資產"].cell(2, 4).quotePrefix is False


def test_characters_that_break_xml_are_removed_instead_of_crashing(store):
    store.add(make_row(name="花瓶\x00\x0b壞字"), None)
    assert store.all_rows()[0].name == "花瓶壞字"


# ------------------------------------------------------------ Excel 被鎖住


def test_locked_file_warns_and_retries_until_it_works(store, monkeypatch):
    import troupe_bot.store as mod

    real_replace = mod.os.replace
    calls = {"n": 0}

    def flaky(src, dst):
        calls["n"] += 1
        if calls["n"] <= 2:
            raise PermissionError("file is open in Excel")
        return real_replace(src, dst)

    monkeypatch.setattr(mod.os, "replace", flaky)
    store.add(make_row(), None)
    assert len(store.warnings) == 2 and "Excel" in store.warnings[0]
    assert len(store.all_rows()) == 1


def test_file_that_stays_locked_raises_excel_busy_and_leaves_no_trace(store, monkeypatch):
    import troupe_bot.store as mod

    store.add(make_row(name="原有"), None)
    before = store.path.read_bytes()
    monkeypatch.setattr(mod.os, "replace", lambda s, d: (_ for _ in ()).throw(PermissionError("locked")))
    with pytest.raises(ExcelBusy):
        store.add(make_row(name="寫不進去"), None)
    assert len(store.warnings) == 3
    assert store.path.read_bytes() == before
    assert not store.path.with_name(store.path.name + ".tmp").exists()


# ------------------------------------------------------------ 更新、刪除、訊息 ID


def test_update_changes_fields_appends_remarks_and_records_who(store):
    store.add(make_row(remarks="借來的"), None)
    row = store.update(1, {"location": "後台A衣箱", "quantity": 4}, by="阿華", at="2026-10-22 16:00:00", remarks_append="要還")
    assert (row.location, row.quantity, row.remarks) == ("後台A衣箱", 4, "借來的；要還")
    assert (row.updated_by, row.updated_at) == ("阿華", "2026-10-22 16:00:00")
    assert store.all_rows()[0].location == "後台A衣箱"


def test_update_of_a_deleted_row_is_refused(store):
    store.add(make_row(), None)
    store.mark_deleted(1, by="阿華", at="2026-10-22 16:00:00")
    with pytest.raises(ValueError):
        store.update(1, {"location": "x"}, by="阿華", at="t")


def test_delete_marks_status_and_keeps_the_row(store):
    store.add(make_row(), None)
    row = store.mark_deleted(1, by="阿華", at="2026-10-22 16:00:00")
    assert row.status == "已刪除" and row.updated_by == "阿華"
    assert len(store.all_rows()) == 1


def test_message_ids_accumulate_and_find_the_row(store):
    store.add(make_row(), None)
    store.add(make_row(name="另一件"), None)
    store.add_message_id(2, 111)
    store.add_message_id(2, 222)
    assert store.find_by_message(222).name == "另一件"
    assert store.find_by_message(111).id == 2
    assert store.find_by_message(999) is None
    assert store.all_rows()[1].message_ids == [111, 222]


def test_unknown_asset_id_raises(store):
    with pytest.raises(KeyError):
        store.mark_deleted(42, by="x", at="t")


# ------------------------------------------------------------ 🆕 新欄位：數量來源、入庫方式、借出人、借出時間

def test_new_columns_round_trip(store):
    store.add(make_row(quantity_ai=True, intake_mode="分開（AI 判定）", borrower="王小明", borrowed_at="2026-10-23 10:00:00"), None)
    [row] = store.all_rows()
    assert (row.quantity_ai, row.intake_mode, row.borrower, row.borrowed_at) == (
        True, "分開（AI 判定）", "王小明", "2026-10-23 10:00:00")
    ws = load_workbook(store.path)["資產"]
    assert ws.cell(2, 9).value == 1 and isinstance(ws.cell(2, 9).value, int)  # 數量維持純整數
    assert ws.cell(2, 19).value == "AI 判斷"


def test_defaults_for_a_plain_row(store):
    store.add(make_row(), None)
    [row] = store.all_rows()
    assert (row.quantity_ai, row.intake_mode, row.borrower, row.borrowed_at) == (False, "一般", "", "")
    assert load_workbook(store.path)["資產"].cell(2, 19).value == "文字"


def test_add_many_writes_all_rows_in_one_save_with_consecutive_ids(store, monkeypatch):
    import troupe_bot.store as mod

    saves = []
    real = mod.os.replace
    monkeypatch.setattr(mod.os, "replace", lambda s, d: (saves.append(1), real(s, d))[1])
    rows = store.add_many([(make_row(name="茶壺"), thumb_bytes()), (make_row(name="扇子"), None), (make_row(name="燈籠"), thumb_bytes())])
    assert [r.id for r in rows] == [1, 2, 3] and len(saves) == 1
    assert [r.name for r in store.all_rows()] == ["茶壺", "扇子", "燈籠"]
    assert len(load_workbook(store.path)["資產"]._images) == 2


def test_add_many_is_all_or_nothing_when_excel_is_locked(store, monkeypatch):
    import troupe_bot.store as mod

    monkeypatch.setattr(mod.os, "replace", lambda s, d: (_ for _ in ()).throw(PermissionError("locked")))
    with pytest.raises(ExcelBusy):
        store.add_many([(make_row(name="甲"), None), (make_row(name="乙"), None)])
    assert store.all_rows() == []


def test_an_old_18_column_file_gets_the_new_headers_and_keeps_its_data(tmp_path):
    from openpyxl import Workbook

    path = tmp_path / "inventory.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "資產"
    ws.append(COLUMNS[:18])
    ws.append([1, "2026-10-22 14:05:09", "小明", "舊的花瓶", "小道具", "青花瓷", "古典", "《白蛇傳》", 2, "良好", "B區03箱",
               "", "", "C:/old.jpg", "小明", "2026-10-22 14:05:09", "正常", "55,66"])
    ws.column_dimensions["R"].hidden = True
    wb.save(path)

    s = Store(path, retries=1, delay=0, sleep=lambda _s: None, warn=lambda m: None)
    s.ensure()
    ws = load_workbook(path)["資產"]
    assert [c.value for c in ws[1]] == COLUMNS
    assert ws.column_dimensions["R"].hidden is True
    [row] = s.all_rows()
    assert (row.name, row.quantity, row.message_ids, row.borrower, row.intake_mode) == ("舊的花瓶", 2, [55, 66], "", "一般")
    s.add(make_row(name="新的"), None)  # 舊檔照樣可以往下加
    assert [r.name for r in s.all_rows()] == ["舊的花瓶", "新的"]


def test_ensure_leaves_an_up_to_date_file_alone(store):
    store.add(make_row(), None)
    before = store.path.read_bytes()
    store.ensure()
    assert store.path.read_bytes() == before


# ------------------------------------------------------------ 🆕 借出 / 歸還

def test_lend_records_borrower_and_time_and_keeps_the_location(store):
    store.add(make_row(), None)
    row = store.lend(1, "王小明", by="阿華", at="2026-10-23 10:00:00")
    assert (row.borrower, row.borrowed_at, row.location) == ("王小明", "2026-10-23 10:00:00", "B區03箱")
    assert (row.updated_by, row.updated_at) == ("阿華", "2026-10-23 10:00:00")
    assert store.all_rows()[0].borrower == "王小明"


def test_lending_an_item_that_is_already_out_is_refused_and_changes_nothing(store):
    store.add(make_row(), None)
    store.lend(1, "王小明", by="阿華", at="2026-10-23 10:00:00")
    before = store.path.read_bytes()
    with pytest.raises(AlreadyLoaned) as err:
        store.lend(1, "李大華", by="阿華", at="2026-10-23 11:00:00")
    assert err.value.row.borrower == "王小明"
    assert store.path.read_bytes() == before


def test_give_back_clears_both_loan_columns(store):
    store.add(make_row(), None)
    store.lend(1, "王小明", by="阿華", at="2026-10-23 10:00:00")
    row = store.give_back(1, by="小美", at="2026-10-24 09:00:00")
    assert (row.borrower, row.borrowed_at, row.updated_by) == ("", "", "小美")
    assert store.all_rows()[0].borrower == ""


def test_returning_an_item_that_is_not_out_is_refused(store):
    store.add(make_row(), None)
    with pytest.raises(NotLoaned):
        store.give_back(1, by="小美", at="t")


def test_a_deleted_item_cannot_be_lent(store):
    store.add(make_row(), None)
    store.mark_deleted(1, by="阿華", at="t")
    with pytest.raises(ValueError):
        store.lend(1, "王小明", by="阿華", at="t")


def test_borrower_names_that_look_like_formulas_are_stored_as_text(store):
    store.add(make_row(), None)
    store.lend(1, '=HYPERLINK("http://x","y")', by="阿華", at="2026-10-23 10:00:00")
    ws = load_workbook(store.path)["資產"]
    cell = ws.cell(2, 21)
    assert cell.data_type == "s" and cell.quotePrefix is True


def test_a_human_correcting_the_quantity_can_clear_the_ai_flag(store):
    store.add(make_row(quantity_ai=True), None)
    row = store.update(1, {"quantity": 6, "quantity_ai": False}, by="阿華", at="t")
    assert (row.quantity, row.quantity_ai) == (6, False)
