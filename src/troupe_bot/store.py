"""Excel 檔（單一工作表「資產」）的讀寫。

- 新的一筆永遠加在最後一列，不覆蓋。
- 每次異動都是「載入 → 修改 → 存成 .tmp → 原子替換」；Excel 正開著時 Windows 會鎖檔，
  替換會失敗，這時印出警告並重試，最後才放棄（ExcelBusy），不留半個壞檔。
- 防護 4（Excel 公式處理）在 _set_text()。
"""

from __future__ import annotations

import io
import logging
import os
import sys
import time
from dataclasses import dataclass, field, fields
from pathlib import Path
from typing import Any, Callable

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font

log = logging.getLogger(__name__)

SHEET = "資產"
COLUMNS = [
    "編號", "登記日期時間", "上傳者", "資產名稱/簡述", "資產類別", "詳細特徵與材質",
    "適用場景/風格", "使用/適用公演", "數量", "物品狀況", "存放位置", "備註",
    "照片縮圖", "照片存放位置", "最後更新者", "更新時間", "狀態", "歸檔訊息ID",
    "數量來源", "入庫方式", "借出人", "借出時間",
]
WIDTHS = [6, 20, 12, 22, 12, 30, 22, 16, 6, 12, 14, 24, 24, 50, 12, 20, 8, 16, 10, 22, 12, 20]
THUMB_COL = 13
PATHS_COL = 14
IDS_COL = 18
QTY_SOURCE_COL = 19
MODE_COL = 20
BORROWER_COL = 21
BORROWED_AT_COL = 22
LAST_COL = 22
QTY_AI_TEXT = "AI 判斷"
QTY_TEXT = "文字"
DEFAULT_MODE = "一般"
THUMB_ROW_HEIGHT = 115
AI_SUFFIX = "（AI）"
STATUS_OK = "正常"
STATUS_DELETED = "已刪除"
FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")
UPDATABLE = {
    "name", "category", "features", "scene", "play", "quantity", "quantity_ai", "condition", "condition_ai", "location",
}


class ExcelBusy(Exception):
    """Excel 檔一直被鎖住（多半是你把它打開了），重試後仍寫不進去。"""


class AlreadyLoaned(Exception):
    """這件已經借出去了，要先歸還才能再借。"""

    def __init__(self, row: "AssetRow"):
        super().__init__(f"已借給 {row.borrower}")
        self.row = row


class NotLoaned(Exception):
    """這件目前沒有借出，不能歸還。"""

    def __init__(self, row: "AssetRow"):
        super().__init__("目前沒有借出")
        self.row = row


def default_warn(message: str) -> None:
    log.warning(message)
    print(f"[警告] {message}", file=sys.stderr, flush=True)


@dataclass(kw_only=True)
class AssetRow:
    id: int = 0
    registered_at: str = ""
    uploader: str = ""
    name: str = ""
    category: str = ""
    features: str = ""
    scene: str = ""
    play: str = ""
    quantity: int = 1
    quantity_ai: bool = False
    condition: str = "良好"
    condition_ai: bool = False
    location: str = "未指定"
    remarks: str = ""
    photo_paths: list[str] = field(default_factory=list)
    updated_by: str = ""
    updated_at: str = ""
    status: str = STATUS_OK
    message_ids: list[int] = field(default_factory=list)
    intake_mode: str = DEFAULT_MODE
    borrower: str = ""
    borrowed_at: str = ""
    row_number: int = field(default=0, repr=False, compare=False)

    @property
    def active(self) -> bool:
        return self.status != STATUS_DELETED

    @property
    def loaned(self) -> bool:
        return bool(self.borrower)


def clean_text(value: Any) -> str:
    """移除 XML 不允許的控制字元（否則 openpyxl 會直接丟例外）。"""
    return ILLEGAL_CHARACTERS_RE.sub("", "" if value is None else str(value))


def _set_text(cell, value: Any) -> None:
    text = clean_text(value)
    cell.value = text  # 不防護的原始值：以 = 開頭的文字會被 openpyxl 當成公式寫入

    # ==== 防護 4：Excel 公式處理　開始 ====
    if text.startswith(FORMULA_PREFIXES):
        cell.data_type = "s"  # 明確存成字串，不是公式
        cell.quotePrefix = True  # 之後在 Excel 裡編輯這格，也不會被重新解析成公式
    # ==== 防護 4 結束 ====


def _write_cells(ws, r: int, row: AssetRow) -> None:
    wrap = Alignment(wrap_text=True, vertical="top")
    cell_values: dict[int, Any] = {
        2: row.registered_at, 3: row.uploader, 4: row.name, 5: row.category, 6: row.features,
        7: row.scene, 8: row.play, 10: row.condition + (AI_SUFFIX if row.condition_ai else ""),
        11: row.location, 12: row.remarks, 15: row.updated_by, 16: row.updated_at, 17: row.status,
        IDS_COL: ",".join(str(i) for i in row.message_ids),
        QTY_SOURCE_COL: QTY_AI_TEXT if row.quantity_ai else QTY_TEXT,
        MODE_COL: row.intake_mode, BORROWER_COL: row.borrower, BORROWED_AT_COL: row.borrowed_at,
    }
    ws.cell(r, 1, row.id)
    ws.cell(r, 9, row.quantity)
    for col, value in cell_values.items():
        _set_text(ws.cell(r, col), value)
    paths_cell = ws.cell(r, PATHS_COL)
    _set_text(paths_cell, "\n".join(row.photo_paths))
    if row.photo_paths and Path(row.photo_paths[0]).is_absolute():
        paths_cell.hyperlink = Path(row.photo_paths[0]).as_uri()
        paths_cell.style = "Hyperlink"
    for col in range(1, LAST_COL + 1):
        ws.cell(r, col).alignment = wrap


def _read_row(ws, r: int) -> AssetRow:
    def text(col: int) -> str:
        value = ws.cell(r, col).value
        return "" if value is None else str(value)

    try:
        quantity = int(ws.cell(r, 9).value)
    except (TypeError, ValueError):
        quantity = 1
    condition = text(10)
    condition_ai = condition.endswith(AI_SUFFIX)
    if condition_ai:
        condition = condition[: -len(AI_SUFFIX)]
    message_ids = [int(p) for p in text(IDS_COL).split(",") if p.strip().isdigit()]
    return AssetRow(
        id=int(ws.cell(r, 1).value),
        registered_at=text(2), uploader=text(3), name=text(4), category=text(5),
        features=text(6), scene=text(7), play=text(8), quantity=quantity,
        condition=condition, condition_ai=condition_ai, location=text(11), remarks=text(12),
        photo_paths=[p for p in text(PATHS_COL).split("\n") if p],
        updated_by=text(15), updated_at=text(16), status=text(17) or STATUS_OK,
        message_ids=message_ids, quantity_ai=text(QTY_SOURCE_COL) == QTY_AI_TEXT,
        intake_mode=text(MODE_COL) or DEFAULT_MODE, borrower=text(BORROWER_COL), borrowed_at=text(BORROWED_AT_COL),
        row_number=r,
    )


def _data_rows(ws):
    for r in range(2, ws.max_row + 1):
        if isinstance(ws.cell(r, 1).value, int):
            yield r


def _find_row(ws, asset_id: int) -> int:
    for r in _data_rows(ws):
        if ws.cell(r, 1).value == asset_id:
            return r
    raise KeyError(asset_id)


class Store:
    def __init__(
        self,
        path: Path,
        *,
        retries: int = 10,
        delay: float = 3.0,
        sleep: Callable[[float], None] = time.sleep,
        warn: Callable[[str], None] = default_warn,
    ):
        self.path = Path(path)
        self.retries = retries
        self.delay = delay
        self.sleep = sleep
        self.warn = warn

    # ---- 建立 / 讀取

    def ensure(self) -> None:
        if self.path.exists():
            self._add_missing_headers()
            return
        self.path.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        ws = wb.active
        ws.title = SHEET
        ws.append(COLUMNS)
        for i, width in enumerate(WIDTHS, start=1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = width
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.column_dimensions["R"].hidden = True
        ws.freeze_panes = "A2"
        wb.save(self.path)

    def _add_missing_headers(self) -> None:
        """舊版的 Excel 檔少了新增的欄位時，在最後面補上表頭；既有資料不動。"""
        if not self._run(lambda ws: any(ws.cell(1, i).value is None for i in range(1, LAST_COL + 1)), save=False):
            return

        def op(ws):
            for i in range(1, LAST_COL + 1):
                cell = ws.cell(1, i)
                if cell.value is None:
                    cell.value = COLUMNS[i - 1]
                    cell.font = Font(bold=True)
                    ws.column_dimensions[cell.column_letter].width = WIDTHS[i - 1]

        self._run(op, save=True)

    def all_rows(self) -> list[AssetRow]:
        return self._run(lambda ws: [_read_row(ws, r) for r in _data_rows(ws)], save=False)

    def find_by_message(self, message_id: int) -> AssetRow | None:
        for row in self.all_rows():
            if message_id in row.message_ids:
                return row
        return None

    # ---- 異動

    def add(self, row: AssetRow, thumb: bytes | None) -> AssetRow:
        return self.add_many([(row, thumb)])[0]

    def add_many(self, items: list[tuple[AssetRow, bytes | None]]) -> list[AssetRow]:
        """一次存檔寫入好幾列（全成功或全不寫），編號連續遞增。"""

        def op(ws):
            ids = [ws.cell(r, 1).value for r in _data_rows(ws)]
            next_id = max(ids, default=0) + 1
            for offset, (row, thumb) in enumerate(items):
                row.id = next_id + offset
                r = ws.max_row + 1
                row.row_number = r
                _write_cells(ws, r, row)
                if thumb:
                    image = XLImage(io.BytesIO(thumb))
                    image.anchor = f"{ws.cell(1, THUMB_COL).column_letter}{r}"
                    ws.add_image(image)
                    ws.row_dimensions[r].height = THUMB_ROW_HEIGHT
            return [row for row, _ in items]

        return self._run(op, save=True)

    def update(self, asset_id: int, changes: dict[str, Any], *, by: str, at: str, remarks_append: str = "") -> AssetRow:
        unknown = set(changes) - UPDATABLE
        if unknown:
            raise ValueError(f"不能更新的欄位：{sorted(unknown)}")

        def op(ws):
            r = _find_row(ws, asset_id)
            row = _read_row(ws, r)
            if not row.active:
                raise ValueError("這筆已經刪除了")
            for key, value in changes.items():
                setattr(row, key, value)
            if remarks_append:
                row.remarks = f"{row.remarks}；{remarks_append}" if row.remarks else remarks_append
            row.updated_by, row.updated_at = by, at
            _write_cells(ws, r, row)
            return row

        return self._run(op, save=True)

    def mark_deleted(self, asset_id: int, *, by: str, at: str) -> AssetRow:
        def op(ws):
            r = _find_row(ws, asset_id)
            row = _read_row(ws, r)
            row.status, row.updated_by, row.updated_at = STATUS_DELETED, by, at
            _write_cells(ws, r, row)
            return row

        return self._run(op, save=True)

    def lend(self, asset_id: int, borrower: str, *, by: str, at: str) -> AssetRow:
        """記下借出人與借出時間；存放位置保留原位置。"""

        def op(ws):
            r = _find_row(ws, asset_id)
            row = _read_row(ws, r)
            if not row.active:
                raise ValueError("這筆已經刪除了")
            if row.loaned:
                raise AlreadyLoaned(row)
            row.borrower, row.borrowed_at = borrower, at
            row.updated_by, row.updated_at = by, at
            _write_cells(ws, r, row)
            return row

        return self._run(op, save=True)

    def give_back(self, asset_id: int, *, by: str, at: str) -> AssetRow:
        def op(ws):
            r = _find_row(ws, asset_id)
            row = _read_row(ws, r)
            if not row.active:
                raise ValueError("這筆已經刪除了")
            if not row.loaned:
                raise NotLoaned(row)
            row.borrower = row.borrowed_at = ""
            row.updated_by, row.updated_at = by, at
            _write_cells(ws, r, row)
            return row

        return self._run(op, save=True)

    def add_message_id(self, asset_id: int, message_id: int) -> None:
        def op(ws):
            r = _find_row(ws, asset_id)
            row = _read_row(ws, r)
            if message_id not in row.message_ids:
                row.message_ids.append(message_id)
                _set_text(ws.cell(r, IDS_COL), ",".join(str(i) for i in row.message_ids))

        self._run(op, save=True)

    # ---- 載入 → 修改 → 原子替換，鎖檔時重試

    def _run(self, fn: Callable[[Any], Any], *, save: bool) -> Any:
        tmp = self.path.with_name(self.path.name + ".tmp")
        last: PermissionError | None = None
        for attempt in range(1, self.retries + 1):
            try:
                wb = load_workbook(self.path)
                result = fn(wb[SHEET])
                if save:
                    try:
                        wb.save(tmp)
                        os.replace(tmp, self.path)
                    finally:
                        try:
                            tmp.unlink(missing_ok=True)
                        except OSError:
                            pass
                return result
            except PermissionError as exc:
                last = exc
                self.warn(
                    f"Excel 檔案正被其他程式使用（可能是你把它打開了），第 {attempt}/{self.retries} 次存取失敗"
                    + (f"，{self.delay:g} 秒後重試…" if attempt < self.retries else "，不再重試。")
                )
                if attempt < self.retries:
                    self.sleep(self.delay)
        raise ExcelBusy(str(self.path)) from last
